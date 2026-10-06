import json
import time
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app as platform
import exiros_fetch
from test_exiros_fetch import CHROMIUM, config_for, needs_chromium, portal  # noqa: F401  (portal is a fixture)
from test_rfq_to_excel import make_report


def report(tmp_path, rfq, price):
    path = tmp_path / f"{rfq}_quotationsReport.xlsx"
    make_report(path, [
        [1, "Dell Monitor", "Dell P2725H monitor", "Tenaris", "2.0", "PZA", "11/30/2026", price, 2 if price else "", 45, "Dell", ""],
        [2, "Cable", "HDMI cable", "Tenaris", "2.0", "PZA", "11/30/2026", "", "", "", "", ""],
    ])
    return path


@pytest.fixture
def client(tmp_path):
    config = tmp_path / "config.json"
    config.write_text("{}")
    return TestClient(platform.create_app(tmp_path / "data", config))


def upload(client, *paths):
    files = [("files", (p.name, p.read_bytes())) for p in paths]
    return client.post("/import", files=files)


def test_urgency_and_time_left():
    now = datetime(2026, 10, 6, 12, 0)
    assert platform.urgency("2026-10-07 10:00", now) == "soon"
    assert platform.urgency("2026-10-11 10:00", now) == "week"
    assert platform.urgency("2026-11-11 10:00", now) == "later"
    assert platform.urgency("2026-10-05 10:00", now) == "closed"
    assert platform.urgency(None, now) == "none"
    assert platform.time_left("2026-10-07 10:00", now) == "22 hours left"
    assert platform.time_left("2026-10-09 12:00", now) == "3 days left"
    assert platform.time_left("2026-10-03 12:00", now) == "Closed 3 days ago"


def test_empty_dashboard_explains_what_to_do(client):
    assert "No RFQs yet" in client.get("/").text


def test_import_view_deadline_and_export(client, tmp_path):
    res = upload(client, report(tmp_path, "5400796", 745), report(tmp_path, "5400801", ""))
    assert "Imported 2 RFQ(s), 2 new" in res.text
    assert res.text.count('class="pill new"') == 2

    soon = (datetime.now() + timedelta(hours=20)).strftime("%Y-%m-%d")
    client.post("/rfq/5400801/deadline", data={"date": soon, "time": "23:00"})
    page = client.get("/rfq/5400801").text
    assert "Not quoted" in page and "left" in page
    dash = client.get("/").text
    assert dash.count('class="pill new"') == 1  # opening 5400801 marked it seen
    assert dash.index("5400801") < dash.index("5400796")  # soonest deadline first
    assert 'alert-soon' in dash

    assert "5400796" in client.get("/?q=monitor").text
    assert "5400796" not in client.get("/?status=Not+quoted").text
    history = client.get("/items?q=dell").text
    assert "745.00" in history and "Times quoted" in history

    res = client.get("/export.xlsx")
    wb = load_workbook(__import__("io").BytesIO(res.content))
    assert wb.sheetnames == ["RFQs", "Lines"]
    rows = list(wb["RFQs"].iter_rows(values_only=True))
    assert [r[0] for r in rows[1:-1]] == ["5400801", "5400796"]
    assert rows[-1][7] == 1490


def test_update_button_without_setup_explains(client):
    res = client.post("/update")
    assert "isn&#39;t set up yet" in res.text or "isn't set up yet" in res.text


@needs_chromium
def test_update_button_runs_full_update(portal, tmp_path, monkeypatch):  # noqa: F811
    base, _ = portal
    config = config_for(base)
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(config))

    def auto_login_browser(p):
        browser = p.chromium.launch(executable_path=CHROMIUM)
        original = browser.new_context

        def new_context(**kwargs):
            context = original(**kwargs)
            context.on("page", lambda page: page.on("load", lambda: page.click("#login") if page.url.endswith("login.html") else None))
            return context

        browser.new_context = new_context
        return browser

    monkeypatch.setattr(exiros_fetch, "open_browser", auto_login_browser)
    client = TestClient(platform.create_app(tmp_path / "data", config_path))
    client.post("/update")
    for _ in range(120):
        state = client.get("/update/status").json()
        if not state["running"]:
            break
        time.sleep(0.5)
    assert state["result"]["ok"], state
    assert "Found 3 RFQs, 3 new" in state["result"]["message"]
    dash = client.get("/?view=all").text
    assert all(n in dash for n in ("5400796", "5400801", "5400815"))
    assert "Last update from Exiros: <strong>just now" in dash
