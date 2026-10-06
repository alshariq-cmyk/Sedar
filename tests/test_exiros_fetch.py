"""Run the downloader against a small fake portal served locally."""

import http.server
import os
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

playwright_api = pytest.importorskip("playwright.sync_api")
import exiros_fetch  # noqa: E402

from test_rfq_to_excel import make_report  # noqa: E402

CHROMIUM = os.environ.get("CHROMIUM_PATH", "/opt/pw-browsers/chromium")
RFQS = ["5400796", "5400801", "5400815"]


def build_portal(root: Path) -> None:
    pages = [RFQS[:2], RFQS[2:]]
    for n, rfqs in enumerate(pages, 1):
        rows = "".join(f'<tr><td><a class="rfq" href="/rfq/{r}.html">RFQ {r}</a></td></tr>' for r in rfqs)
        nxt = f'<a id="next" href="/list{n + 1}.html">Next</a>' if n < len(pages) else ""
        (root / f"list{n}.html").write_text(f"<html><body><table>{rows}</table>{nxt}</body></html>")
    (root / "rfq").mkdir()
    (root / "files").mkdir()
    for r in RFQS:
        (root / "rfq" / f"{r}.html").write_text(
            f'<html><body><h1>{r}</h1><a id="export" href="/files/{r}.xlsx" download="quotationsReport.xlsx">Export</a></body></html>')
        make_report(root / "files" / f"{r}.xlsx",
                    [[1, "Item", "Item", "Tenaris", "2.0", "PZA", "11/30/2026", 100, 2, 30, "", ""]])


@pytest.fixture
def portal(tmp_path):
    site = tmp_path / "site"
    site.mkdir()
    build_portal(site)
    handler = lambda *a, **k: http.server.SimpleHTTPRequestHandler(*a, directory=str(site), **k)  # noqa: E731
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    handler_cls = server.RequestHandlerClass
    handler_cls.log_message = lambda *a: None
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


@pytest.mark.skipif(not Path(CHROMIUM).exists(), reason="no local Chromium")
def test_run_downloads_every_rfq_and_builds_excel(portal, tmp_path, monkeypatch):
    monkeypatch.setattr(exiros_fetch, "DOWNLOAD_DIR", tmp_path / "downloads")
    monkeypatch.setattr(exiros_fetch, "OUTPUT_DIR", tmp_path / "output")
    config = {
        "rfq_list_url": f"{portal}/list1.html", "rfq_link": "a.rfq", "rfq_number_regex": r"\d{6,}",
        "next_page": "#next", "download_button": "#export", "delay_seconds": 0,
    }
    with playwright_api.sync_playwright() as p:
        browser = p.chromium.launch(executable_path=CHROMIUM)
        page = browser.new_context(accept_downloads=True).new_page()
        output = exiros_fetch.run(page, config)
        assert sorted(f.name for f in (tmp_path / "downloads").iterdir()) == [f"{r}_quotationsReport.xlsx" for r in RFQS]

        # A second run skips what is already downloaded.
        (tmp_path / "downloads" / f"{RFQS[0]}_quotationsReport.xlsx").write_bytes(
            (tmp_path / "downloads" / f"{RFQS[1]}_quotationsReport.xlsx").read_bytes())
        exiros_fetch.run(page, config)
        browser.close()

    summary = list(load_workbook(output)["Summary"].iter_rows(values_only=True))
    assert [r[0] for r in summary[1:-1]] == RFQS
    assert summary[-1][5] == 600


def test_run_refuses_until_config_filled(tmp_path):
    with pytest.raises(SystemExit, match="isn't filled in"):
        exiros_fetch.run(page=None, config={"rfq_link": "", "download_button": ""})
