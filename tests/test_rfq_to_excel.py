import pandas as pd
from openpyxl import load_workbook

from rfq_to_excel import main, read_report, rfq_number


def make_report(path, rows):
    """Mimic the Exiros layout: blank first row/column, header on row 2."""
    header = ["", " # ", "Description", "Long Description", "Client", "Requested quantity", "Unit of measure",
              "Date needed", "Unit price", "Quoted quantity", "Delivery time", "Brand", "Notes"]
    pd.DataFrame([[""] * len(header), header] + [[""] + r for r in rows]).to_excel(path, header=False, index=False)


def test_reads_exiros_layout_and_takes_rfq_from_filename(tmp_path):
    path = tmp_path / "ab12-5400796_quotationsReport.xlsx"
    make_report(path, [
        [1, "Monitor", "Monitor long", "Tenaris", "2.0", "PZA - Piece", "01/04/2027", 745, 2, 60, "Dell", ""],
        [2, "Envelopes", "Env long", "Tenaris", "500.0", "PZA - Piece", "11/30/2026", 1.8, 400, 45, "", "partial"],
        [3, "Cable", "Cable long", "Tenaris", "1.0", "PZA - Piece", "11/30/2026", "", "", "", "", ""],
    ])
    assert rfq_number(path) == "5400796"
    df = read_report(path)
    assert df["rfq"].unique().tolist() == ["5400796"]
    assert df["line_total"].tolist()[:2] == [1490, 720]
    assert df["status"].tolist() == ["Quoted", "Qty differs", "Not quoted"]
    assert df["date_needed"].iloc[0].strftime("%Y-%m-%d") == "2027-01-04"


def test_combines_folder_into_one_workbook(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    for rfq in ("5400796", "5400801"):
        make_report(folder / f"{rfq}_quotationsReport.xlsx",
                    [[1, "Item", "Item", "Tenaris", "1.0", "PZA", "11/30/2026", 100, 1, 30, "", ""]])
    out = tmp_path / "all.xlsx"
    assert main([str(out), str(folder)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Lines"]
    summary = [row for row in wb["Summary"].iter_rows(values_only=True)]
    assert [r[0] for r in summary] == ["RFQ No.", "5400796", "5400801", "Total"]
    assert summary[-1][5] == 200
