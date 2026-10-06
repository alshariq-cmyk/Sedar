"""Combine Exiros "quotationsReport" exports into one formatted Excel workbook.

    python rfq_to_excel.py OUTPUT.xlsx REPORT.xls [REPORT.xls ...]
    python rfq_to_excel.py OUTPUT.xlsx path/to/folder

Each export from the Exiros supplier portal covers one RFQ and has an "Items"
sheet with the requested lines and our submitted price, quantity, delivery
time, brand and notes. The RFQ number is only in the file name
(e.g. 5400796_quotationsReport.xls), so it is taken from there.

The output has a "Summary" sheet (one row per RFQ) and a "Lines" sheet (every
line of every RFQ, with line totals and a status column).
"""

import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Exiros header -> our column name
COLUMNS = {
    "#": "line",
    "Description": "description",
    "Long Description": "long_description",
    "Client": "client",
    "Requested quantity": "requested_qty",
    "Unit of measure": "uom",
    "Date needed": "date_needed",
    "Unit price": "unit_price",
    "Quoted quantity": "quoted_qty",
    "Delivery time": "delivery_days",
    "Brand": "brand",
    "Notes": "notes",
}
REQUIRED = {"#", "Description", "Unit price"}


def rfq_number(path: Path) -> str:
    match = re.search(r"(\d{5,})_quotationsReport", path.name, re.IGNORECASE)
    return match.group(1) if match else path.stem


def read_report(path: Path) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    header_row = next(
        (i for i in range(min(20, len(raw))) if REQUIRED <= {str(v).strip() for v in raw.iloc[i] if pd.notna(v)}),
        None,
    )
    if header_row is None:
        raise ValueError(f"{path.name}: couldn't find the Exiros item header row (#, Description, Unit price…)")
    headers = [str(v).strip() if pd.notna(v) else "" for v in raw.iloc[header_row]]
    return lines_from_table(headers, raw.iloc[header_row + 1 :].values.tolist(), rfq_number(path), path.name)


def lines_from_table(headers: list[str], rows: list[list], rfq: str, source: str) -> pd.DataFrame:
    """Normalize item rows laid out like the Exiros items table (from the export file or the quote page)."""
    headers = [re.sub(r"\s+", " ", str(h)).strip() for h in headers]
    df = pd.DataFrame([list(r) + [None] * (len(headers) - len(r)) for r in rows], columns=headers, dtype=object)
    df = df.loc[:, ~df.columns.duplicated()]
    df = df[[h for h in df.columns if h in COLUMNS]].rename(columns=COLUMNS)
    for name in COLUMNS.values():
        if name not in df:
            df[name] = None
    df = df[df["line"].notna() & (df["line"].astype(str).str.strip() != "")]

    for name in ("line", "requested_qty", "unit_price", "quoted_qty", "delivery_days"):
        # The quote page shows numbers like "3,290.0".
        df[name] = pd.to_numeric(df[name].map(lambda v: str(v).replace(",", "").strip() if isinstance(v, str) else v), errors="coerce")
    df["line"] = df["line"].astype("Int64")
    # The portal writes dates as MM/DD/YYYY text.
    df["date_needed"] = pd.to_datetime(df["date_needed"], format="%m/%d/%Y", errors="coerce")
    for name in ("description", "long_description", "client", "uom", "brand", "notes"):
        df[name] = df[name].map(lambda v: str(v).strip() if pd.notna(v) and str(v).strip() else None)

    df["line_total"] = df["unit_price"] * df["quoted_qty"].fillna(df["requested_qty"])

    def status(r) -> str:
        if pd.isna(r["unit_price"]) or r["unit_price"] == 0:
            return "Not quoted"
        if pd.notna(r["quoted_qty"]) and pd.notna(r["requested_qty"]) and r["quoted_qty"] != r["requested_qty"]:
            return "Qty differs"
        return "Quoted"

    df["status"] = df.apply(status, axis=1) if len(df) else pd.Series(dtype=object)
    df.insert(0, "rfq", rfq)
    df["source_file"] = source
    return df.reset_index(drop=True)


def summarize(lines: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for rfq, g in lines.groupby("rfq", sort=False):
        rows.append({
            "rfq": rfq,
            "client": "; ".join(dict.fromkeys(c for c in g["client"] if c)),
            "lines": len(g),
            "quoted": int((g["status"] != "Not quoted").sum()),
            "not_quoted": int((g["status"] == "Not quoted").sum()),
            "total": float(g["line_total"].sum()),
            "earliest_needed": g["date_needed"].min(),
            "max_delivery_days": g["delivery_days"].max(),
            "source_file": g["source_file"].iloc[0],
        })
    return pd.DataFrame(rows)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {"Not quoted": PatternFill("solid", fgColor="F8D7DA"), "Qty differs": PatternFill("solid", fgColor="FFF3CD")}
MONEY = "#,##0.00"
QTY = "#,##0.##"
DATE = "dd-mmm-yyyy"


def write_sheet(ws, frame: pd.DataFrame, spec: list[tuple[str, str, int, str | None]], total_cols: tuple[str, ...] = ()):
    """spec: (column, header, width, number_format)."""
    for c, (_, header, width, _) in enumerate(spec, 1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 32
    for r, record in enumerate(frame.to_dict("records"), 2):
        for c, (name, _, _, fmt) in enumerate(spec, 1):
            value = record.get(name)
            if value is None or (not isinstance(value, str) and pd.isna(value)):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=name in ("long_description", "client", "notes", "description"))
            if fmt:
                cell.number_format = fmt
            if name == "status" and value in STATUS_FILL:
                cell.fill = STATUS_FILL[value]
    last = len(frame) + 1
    if total_cols and len(frame):
        ws.cell(row=last + 1, column=1, value="Total").font = TOTAL_FONT
        for c, (name, _, _, fmt) in enumerate(spec, 1):
            if name in total_cols:
                col = get_column_letter(c)
                cell = ws.cell(row=last + 1, column=c, value=float(frame[name].fillna(0).sum()))
                cell.font, cell.number_format = TOTAL_FONT, fmt or MONEY
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec))}{last}"


def build_workbook(lines: pd.DataFrame, output: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_sheet(summary, summarize(lines), [
        ("rfq", "RFQ No.", 12, None),
        ("client", "Client", 45, None),
        ("lines", "Lines", 8, "0"),
        ("quoted", "Quoted", 9, "0"),
        ("not_quoted", "Not quoted", 10, "0"),
        ("total", "Total quoted value", 18, MONEY),
        ("earliest_needed", "Earliest date needed", 16, DATE),
        ("max_delivery_days", "Longest delivery (days)", 14, "0"),
        ("source_file", "Source file", 32, None),
    ], total_cols=("lines", "quoted", "not_quoted", "total"))

    ws = wb.create_sheet("Lines")
    write_sheet(ws, lines, [
        ("rfq", "RFQ No.", 12, None),
        ("line", "#", 5, "0"),
        ("description", "Description", 32, None),
        ("long_description", "Long description", 55, None),
        ("client", "Client", 36, None),
        ("requested_qty", "Requested qty", 11, QTY),
        ("uom", "UoM", 12, None),
        ("date_needed", "Date needed", 13, DATE),
        ("unit_price", "Our unit price", 13, MONEY),
        ("quoted_qty", "Quoted qty", 10, QTY),
        ("line_total", "Line total", 14, MONEY),
        ("delivery_days", "Delivery (days)", 10, "0"),
        ("brand", "Brand", 18, None),
        ("notes", "Notes", 30, None),
        ("status", "Status", 12, None),
    ], total_cols=("line_total",))
    wb.properties.title = "Exiros RFQ quotations"
    wb.properties.created = datetime.now()
    wb.save(output)


def collect(paths: list[Path]) -> list[Path]:
    files: list[Path] = []
    for p in paths:
        if p.is_dir():
            files += sorted(f for f in p.rglob("*") if f.suffix.lower() in (".xls", ".xlsx") and not f.name.startswith("~$"))
        else:
            files.append(p)
    return files


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    output, inputs = Path(argv[0]), collect([Path(a) for a in argv[1:]])
    frames, failed = [], []
    for f in inputs:
        try:
            frames.append(read_report(f))
        except Exception as exc:  # keep going so one bad file doesn't block the rest
            failed.append(f"{f.name}: {exc}")
    if not frames:
        print("No RFQ files could be read.", *failed, sep="\n")
        return 1
    lines = pd.concat(frames, ignore_index=True)
    build_workbook(lines, output)
    print(f"Wrote {output}: {len(frames)} RFQ(s), {len(lines)} line(s), total {lines['line_total'].sum():,.2f}")
    for msg in failed:
        print("Skipped", msg)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
